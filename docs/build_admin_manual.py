from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / 'manual-assets'
ASSETS.mkdir(exist_ok=True)
OUT = ROOT / 'Lightmare_PH_Admin_Visual_Manual.docx'

# Images come from an isolated rendering of the unmodified admin components.
PAGES = [
('Lightmare PH Admin Visual Manual', '/admin', 'Overview',
 ['Overview | Store activity and email allowance', 'Orders | Find and process customer orders', 'Products | Create and maintain the catalog', 'Landing Page | Edit and publish storefront content'],
 'Use this manual to run the Lightmare PH back office, from reviewing orders to updating the storefront. Follow the numbered steps and check the result before moving to the next task.',
 [('Start here', 'Sign in on page 2 and review Overview on page 3. Pages 4 to 6 cover orders; pages 7 to 9 cover products. Pages 10 to 12 cover the landing page. Payments and shipping are on pages 13 and 14. Account settings and troubleshooting are on page 15; the daily checklist is on page 16.')],
 'Screenshots show the current admin interface with sample data. Customer names, orders, totals, and payment instructions are examples. Use your approved store information when entering data.'),
('Sign in and find your way', '/admin/login', None,
 ['Email | Your administrator email', 'Password | Your current password', 'Sign in | Open the back office', 'Back to Lightmare PH | Return to the storefront'],
 'Open the store website followed by /admin/login. Have your administrator email and current password ready.',
 [('1 Enter your credentials', 'Fill in Email and Password, then select Sign in. A successful sign in opens Overview.'), ('2 Use the navigation', 'Use Overview, Orders, Products, Landing Page, Settings, Payments, and Shipping in the sidebar. On a small screen, use the navigation menu button.'), ('3 Check customer views', 'View storefront opens the home page. Browse products opens the shop. Use these links to check published changes.'), ('4 Finish your session', 'Select Sign out when you finish. The sign in page should appear. If sign out fails, retry and check the message.')],
 'If a protected screen returns you to sign in, sign in again and reopen the task. Do not assume an unsaved form survived.'),
('Review the overview', '/admin', 'Overview',
 ['Total products | Published products and drafts', 'Pending payment | Orders awaiting payment', 'Email needs attention | Delivery issues to review', 'Daily emails used Brevo | Daily allowance and remaining credits'],
 'Use Overview at the start of each work session to identify orders and email issues that need attention.',
 [('1 Read the totals', 'Total products includes drafts and published products, excluding deleted products. Total orders covers all orders. Paid order value includes shipping.'), ('2 Review pending work', 'Open Orders to inspect pending orders. Email needs attention indicates order email records that need review.'), ('3 Check the daily allowance', 'Read the Brevo allowance and the remaining batch estimate. Resends use the allowance. Refresh Overview to check current credits.'), ('4 Open recent records', 'Select a recent order reference or product name to find it in the appropriate list. Overview shows the five latest orders and products; totals cover all records.')],
 'Unavailable or Not configured is not a zero balance. If credits or totals cannot load, refresh once and use the troubleshooting steps on page 15.'),
('Find and inspect an order', '/admin/orders', 'Orders',
 ['Search orders | Name email or order reference', 'Filter by status | All statuses or a specific status', 'Order row | Reference customer items total status email', 'Details | Open the complete order'],
 'Find the correct order before taking any action. The order reference is the clearest identifier when a customer has placed several orders.',
 [('1 Search', 'Enter the customer name, email, or order reference in Search orders. Use the status filter to narrow the list. Use pagination controls when the results span several pages.'), ('2 Open Details', 'Select Details on the matching row. Confirm the customer name, email, phone, delivery address, and reference.'), ('3 Check the contents', 'Review every product, size, color, quantity, unit price, and subtotal. Then check shipping, total, payment method, and notes.'), ('4 Review payment instructions', 'If shown, Instructions at order time records the instructions saved with that order. Use Back to orders to return to the list.')],
 'Changing a product or current payment instructions does not rewrite the details captured for an existing order.'),
('Update order status and resend emails', '/admin/orders', 'Orders',
 ['Status for the entire order | pending paid processing shipped cancelled', 'Resend order emails | Retry customer and configured admin emails', 'Order message | Check the success or failure notice', 'Back to orders | Review the updated list'],
 'Order status applies to the entire order. Selecting a different status starts the update immediately; there is no separate status save button.',
 [('1 Verify the next status', 'Use pending while payment is outstanding. Mark paid only after verifying payment separately. Use processing for preparation, shipped after dispatch, or cancelled when the order is cancelled under your store procedure.'), ('2 Change the status', 'Select the required value under Status for the entire order. Wait for Order status updated. Confirm the selected status and the list badge.'), ('3 Retry delivery when needed', 'Select Resend order emails only when a retry is needed. This resends the customer invoice and configured admin summary using the current status, including emails already delivered.'), ('4 Check the result', 'Wait for Order emails sent or read the error. Review the Email column and Overview allowance before repeated retries.')],
 'A status update confirms the order record change. Use the explicit resend action when you need to send the current status to the email recipients.'),
('Adjust shipping on a pending order', '/admin/orders', 'Orders',
 ['Adjust shipping | Available for pending orders', 'New shipping fee | Enter the agreed fee in pesos', 'Reason | Explain why the fee is changing', 'Save shipping and send receipt | Update total and send revised receipt'],
 'Use this task for an individual pending order when an agreed shipping fee differs from the saved fee. General shipping settings are covered on page 14.',
 [('1 Confirm the order', 'Open Details and check that the order is pending. The adjustment form is shown only for pending orders.'), ('2 Enter the fee and reason', 'Enter New shipping fee in pesos, including decimals if needed. Use a clear reason such as an agreed courier fee. The reason must contain at least three characters.'), ('3 Save once', 'Select Save shipping and send receipt. The order total is recalculated and a revised receipt is sent to the customer.'), ('4 Verify both outcomes', 'Check the new shipping fee, total, and Shipping adjustment history. Read the confirmation to see whether email delivery was confirmed. If it was not, use Resend order emails to retry.')],
 'Example: PHP 690.00 in products plus PHP 150.00 shipping gives PHP 840.00 total. If another administrator changes the order first, reload and review the current values before retrying.'),
('Find and manage products', '/admin/products', 'Products',
 ['Search products | Find a catalog entry', 'Add product | Create a new entry', 'Edit | Change an existing product', 'Delete | Open the deletion confirmation'],
 'Use Products to maintain the catalog shown in the storefront and Overview. A draft is hidden from the shop.',
 [('1 Find the product', 'Search the product list and use pagination as needed. Check the name, price, and visibility before opening Edit.'), ('2 Create or edit', 'Select Add product to start a new entry, or Edit beside an existing one. Complete the fields described on page 8.'), ('3 Save and check', 'Select Save product and wait for Product saved. Storefront and overview now use the updated catalog. Check the product list and Browse products.'), ('4 Hide or remove', 'To hide a product, edit its Visibility to Draft and save. Use Delete only when removal from the catalog is intended; see page 9.')],
 'This screen manages available sizes and visibility. It does not provide a stock quantity field; keep your operational stock records current separately.'),
('Create or edit a product', '/admin/products', 'Products',
 ['Product name and URL slug | Example Tee and example-tee', 'Description and Price PHP | Customer copy and selling price', 'Color name and Available sizes | Vintage white and S M L XL', 'Visibility and Save product | Draft or Published then save'],
 'Prepare the approved product name, description, price, color, available sizes, and image before entering the editor.',
 [('1 Enter identity and copy', 'Fill in Product name and Description. A new product generates its URL slug from the name. Custom slugs use lowercase letters, numbers, and hyphens. Editing an existing name keeps its slug; editing the slug changes shared links.'), ('2 Set the price and color', 'Enter Price (PHP) in pesos, for example 690.00. Fill in Color name using the name customers should see.'), ('3 Select sizes and appearance', 'Choose at least one Available size from S, M, L, and XL. Choose Visibility and Card background. New entries start with Published selected; change this to Draft if they are not ready.'), ('4 Save', 'Add the image using page 9, then select Save product. Verify the success notice and customer product page.')],
 'Suggested preparation: create as Draft, check the saved details, then switch to Published and save when the item is ready for sale.'),
('Add images and control visibility', '/admin/products', 'Products',
 ['Upload product image | JPEG PNG or WebP up to 10 MB', 'Product image URL | Hosted HTTPS image or local image path', 'Visibility | Published or Draft hidden from shop', 'Delete product | Confirm removal or select Cancel'],
 'Use a clear product image and check its appearance before publishing. Image upload and product save are separate actions.',
 [('1 Upload or supply an image', 'Choose a JPEG, PNG, or WebP file up to 10 MB. Wait for Image ready and check the selected preview. You can also enter a hosted HTTPS image URL or an existing /images/ path.'), ('2 Apply the image', 'Select Save product to apply the uploaded image or URL. An empty image displays Image coming soon.'), ('3 Publish or hide', 'Choose Published to show the product in the shop, or Draft (hidden from shop) to hide it. Save and use Browse products to verify.'), ('4 Delete carefully', 'Select Delete in the list and check the named product in the confirmation. Select Delete product to remove it from the catalog and prevent new orders, or Cancel to keep it. Existing order records remain.')],
 'If upload fails, confirm the file format and size first. Do not treat a successful upload as a successful product save.'),
('Edit the landing page and hero', '/admin/content', 'Landing Page',
 ['Page content | Hero and page sections', 'Hero image heading and button | Upload image and edit overlay', 'Unpublished changes | Changes are still in the editor', 'Publish changes | Apply the edited landing page'],
 'Landing Page has four categories: Page content, Shopping help, Contact & socials, and Settings. Open the relevant section to edit its fields.',
 [('1 Edit the hero', 'Under Page content, open Hero image, heading & button. Upload an image or edit its URL. Review alternative text, heading, subtitle, button text, and button destination.'), ('2 Check the composition', 'Adjust text color, button background, button text color, and Horizontal image crop. Mobile uses a 4:5 image crop. Heading and subtitle line breaks are preserved.'), ('3 Edit other page sections', 'Open the relevant collection, about, or bottom call to action section. Use accurate customer copy and verify button destinations.'), ('4 Publish and review', 'Wait for uploads to finish. Select Publish changes, then check the Published confirmation. Use View published landing page to check the result on desktop and mobile.')],
 'An upload alone does not update the storefront. Discard changes restores the last saved content in the editor. Preserve needed edits before navigating away or reloading.'),
('Update shopping help and FAQs', '/admin/content', 'Landing Page',
 ['Shopping help | Size guide ordering shipping and FAQs', 'Order steps | Edit each title and description', 'Size measurements cm | Width and length for each size', 'FAQ questions and answers | Add FAQ Move up or Remove'],
 'Keep shopping help aligned with the products and policies customers can actually use.',
 [('1 Update instructions', 'Choose Shopping help. Edit ordering headings and the Order steps titles and descriptions. Update the shipping coverage description and delivery estimate as needed.'), ('2 Check measurements', 'Open Size measurements (cm). Enter width and length using confirmed garment measurements. Product size availability is managed separately in Products.'), ('3 Maintain FAQs', 'Open FAQ questions & answers. Edit each question and answer. Use Add FAQ + to add an entry, Move up to reorder it, or Remove to delete it from the editor. The editor supports one to twelve FAQs.'), ('4 Publish and inspect', 'Select Publish changes. Check the published size guide, ordering instructions, shipping text, and FAQs for readable and accurate information.')],
 'Shipping help describes delivery to customers. Actual checkout fees and the free shipping threshold are controlled in Shipping, not by changing this text.'),
('Update contact links and page settings', '/admin/content', 'Landing Page',
 ['Contact and socials | Public contact email Instagram URL TikTok URL', 'Settings | Browser tab icon SEO Header and footer', 'Browser tab icon | Upload a square image then publish', 'Page title and Search description | Review customer facing metadata'],
 'These controls maintain the public contact details, links, browser tab icon, and landing page metadata.',
 [('1 Set public contact details', 'Choose Contact & socials. Enter Public contact email and the full Instagram and TikTok URLs. Confirm that each link points to the intended account.'), ('2 Update the tab icon', 'Choose Settings and open Browser tab icon. Upload a square PNG, JPEG, or WebP up to 10 MB. Check the preview or select Restore default icon if appropriate.'), ('3 Review page copy', 'Open SEO to edit Page title and Search description. Open Header & footer to edit the announcement and available copy.'), ('4 Publish and verify', 'Select Publish changes. Check the storefront contact links, announcement, browser tab title, and icon. Refresh the customer page if needed.')],
 'Public contact email is separate from Brevo sender and admin notification settings. Editing it does not configure outgoing email delivery.'),
('Configure payment instructions', '/admin/payments', 'Payments',
 ['Offer GCash | Enable and enter approved instructions', 'Offer bank transfer | Enable and enter approved instructions', 'Offer cash on delivery COD | No advance payment instructions', 'Publish payment settings | Save or Discard changes'],
 'Payments controls the methods and instructions offered to customers. The website does not collect or verify payments through a payment gateway.',
 [('1 Choose methods', 'Enable Offer GCash, Offer bank transfer, or Offer cash on delivery (COD) according to your approved store policy.'), ('2 Enter clear instructions', 'For GCash and bank transfer, provide the approved recipient details and customer instructions. Check account names and numbers against your trusted records before publishing.'), ('3 Publish', 'Select Publish payment settings and wait for Payment instructions published. Use Discard changes if you want to restore the last saved settings.'), ('4 Check checkout', 'Review the payment choices and instructions in the customer checkout without placing an order. COD displays no advance payment instructions.')],
 'Existing orders retain their saved payment instructions. Payment confirmation remains an administrator task; verify separately before marking an order paid.'),
('Configure shipping fees', '/admin/shipping', 'Shipping',
 ['Delivery areas PHP | Metro Manila Rest of Luzon Visayas Mindanao', 'Free shipping | Enable a minimum product subtotal', 'Minimum product subtotal PHP | Products only excluding shipping', 'Publish shipping settings | Apply rates to new orders'],
 'Shipping settings determine checkout shipping fees for new orders. Existing orders keep their saved shipping fee.',
 [('1 Set delivery area fees', 'Enter approved fees in pesos for Metro Manila, Rest of Luzon, Visayas, and Mindanao. Use decimals for centavos if needed.'), ('2 Set free shipping', 'Enable Offer free shipping above a minimum product subtotal if the store offers it. Set Minimum product subtotal (PHP). Customers at or above the threshold receive free shipping in any supported delivery area.'), ('3 Publish', 'Select Publish shipping settings and wait for Shipping rates published. Discard changes restores the saved fees in the editor.'), ('4 Check customer totals', 'Review checkout below and at the threshold, and for the supported delivery areas. Confirm the displayed fee and total without submitting an order.')],
 'The threshold uses the product subtotal only. For one pending order that needs a different fee, use Adjust shipping in its Details panel, as described on page 6.'),
('Account settings and troubleshooting', '/admin/settings', 'Settings',
 ['Current password | Enter your existing password', 'New password | At least 12 characters', 'Confirm new password | Repeat the new password', 'Change password | Other admin sessions are signed out'],
 'Use Settings when you need to change the administrator password. Read any error before retrying an action elsewhere in the back office.',
 [('Change the password', 'Enter Current password, New password, and Confirm new password. Use at least 12 characters. Select Change password and wait for the confirmation. Every other admin session is signed out.'), ('Another session changed settings', 'Before reloading, copy any text you need to keep. Reload the editor, review the latest content or settings, reapply your edits, and publish again. Unsaved changes are not committed after a conflict.'), ('Save or upload fails', 'Check the highlighted field and message. Image uploads require JPEG, PNG, or WebP up to 10 MB. Products require at least one size. If the message says Sign in required, sign in again.'), ('Orders or configuration cannot load', 'Refresh once. If the error persists, give your technical administrator the screen name, action, and exact error. Database, email, and setup configuration may need attention.')],
 'For unconfirmed email delivery, check Brevo allowance and retry using Resend order emails when appropriate. Never assume a failed email message means the order record was not updated.'),
('Daily checklist and publishing checks', '/admin', 'Overview',
 ['Start | Review Overview and daily email allowance', 'Work | Verify orders and maintain catalog details', 'Publish | Save changes and check customer views', 'Finish | Review outstanding issues and Sign out'],
 'Use this checklist to finish each work session with clear records and verified customer information.',
 [('At the start', 'Review Overview totals, pending payment orders, and Email needs attention. Check the remaining daily email allowance before sending retries.'), ('During order work', 'Confirm the reference, customer, items, address, and totals. Verify payment separately. Apply the correct whole order status. Check both record and email results after a shipping adjustment.'), ('Before publishing', 'Check product prices, images, sizes, slugs, and visibility. Review hero crop and contrast, customer help, FAQ answers, contact links, payment recipient details, shipping fees, and the free shipping threshold.'), ('Before finishing', 'Confirm each success notice. Inspect the storefront, shop, product page, and affected checkout fields. Record unresolved errors for follow up. Use Sign out and confirm that the sign in page appears.')],
 'Quick reference: order search and details page 4; status and emails page 5; shipping adjustment page 6; products pages 7 to 9; landing page pages 10 to 12; payments page 13; shipping page 14; troubleshooting page 15.')
]

SCREENS = [
 [('overview', 'Overview and admin navigation')],
 [('login', 'Administrator sign in')],
 [('overview', 'Overview totals and email allowance')],
 [('orders', 'Order search filters and Details buttons')],
 [('order-actions', 'Whole order status and Resend order emails')],
 [('shipping-adjustment', 'Shipping adjustment fee reason and save action')],
 [('products', 'Product list and Add product Edit and Delete actions')],
 [('product-editor', 'Edit product identity description and price')],
 [('product-save-controls', 'Available sizes visibility and Save product'), ('product-delete', 'Product deletion confirmation')],
 [('hero-controls', 'Hero text colors and horizontal image crop controls')],
 [('faq-entry', 'FAQ question answer and ordering actions')],
 [('contact-form', 'Public contact email Instagram URL and TikTok URL')],
 [('payments-form', 'Payment methods instructions and publishing actions')],
 [('shipping-form', 'Delivery fees free shipping threshold and publishing actions')],
 [('settings', 'Change password form')],
 [('overview', 'Overview for the daily review')],
]

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = Inches(.62)
sec.bottom_margin = Inches(.62)
sec.left_margin = sec.right_margin = Inches(.67)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Caption']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    for border in style._element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    for element in style._element.xpath('.//w:rFonts'):
        for key in list(element.attrib):
            if 'theme' in key.lower():
                del element.attrib[key]
    style.paragraph_format.space_after = Pt(7)
doc.styles['Normal'].font.size = Pt(10)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(25)
doc.styles['Heading 1'].font.size = Pt(21)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(9)
doc.styles['Caption'].font.size = Pt(8)
doc.styles['Caption'].font.italic = False
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Lightmare PH Admin Manual  |  ').font.size = Pt(8)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = 'Lightmare PH Admin Visual Manual'
doc.core_properties.subject = 'Illustrated operating guide for store administrators'
doc.core_properties.author = 'Lightmare PH'
for i, (title, route, active, rows, intro, steps, note) in enumerate(PAGES, 1):
    if i > 1:
        doc.add_page_break()
    doc.add_paragraph(title, 'Title' if i == 1 else 'Heading 1')
    p = doc.add_paragraph(f'{i:02}  /  16     {route}')
    p.runs[0].font.size = Pt(9)
    doc.add_paragraph(intro)
    figures = SCREENS[i-1]
    for name, caption in figures:
        path = ASSETS / 'screenshots' / (name + '.png')
        w, h = Image.open(path).size
        maximum_height = 3.50 / len(figures)
        width = min(6.90, maximum_height * w / h)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        inline = p.add_run().add_picture(str(path), width=Inches(width))
        inline._inline.docPr.set('descr', 'Screenshot of the actual admin interface with sample data. ' + caption)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        p = doc.add_paragraph(caption, 'Caption')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for heading, body in steps:
        doc.add_paragraph(heading, 'Heading 2')
        doc.add_paragraph(body)
    p = doc.add_paragraph()
    p.add_run('Remember  ').bold = True
    p.add_run(note)
for p in doc.paragraphs:
    if p.style.name == 'Normal':
        p.style = doc.styles['Normal']
    for run in p.runs:
        run.font.name = 'Arial'
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.size = Pt({'Title':25,'Heading 1':21,'Heading 2':12,'Caption':8}.get(p.style.name,10))
doc.save(OUT)
print(OUT)
print(f'{len(PAGES)} sections and {len(doc.inline_shapes)} embedded screenshots')
